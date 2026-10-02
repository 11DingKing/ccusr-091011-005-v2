"""
仓库管理视图
"""
import logging
import io
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval, Seal
from . import lineage
from .lineage import LineageError
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    SealSerializer, SealCreateSerializer, SplitSerializer, MergeSerializer,
    RepackSerializer, IssueSealsSerializer, FreezeSealsSerializer,
)

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 封签谱系 ====================

def _first_error(errors):
    """从嵌套的序列化器错误中取出第一条人类可读信息。"""
    if isinstance(errors, dict):
        for value in errors.values():
            message = _first_error(value)
            if message:
                return message
    elif isinstance(errors, list):
        for item in errors:
            if isinstance(item, str):
                return str(item)
            message = _first_error(item)
            if message:
                return message
    return '参数校验失败'


def _validate(serializer):
    """校验入参，失败时返回 (None, 错误响应)。"""
    if not serializer.is_valid():
        return None, error_response(message=_first_error(serializer.errors))
    return serializer.validated_data, None


def _load_seals(seal_nos):
    """按封签号批量加载，缺失时返回 None 与错误信息。"""
    seals = list(Seal.objects.filter(seal_no__in=seal_nos))
    if len(seals) != len(set(seal_nos)):
        found = {s.seal_no for s in seals}
        missing = [no for no in seal_nos if no not in found]
        return None, f'封签不存在：{"、".join(missing)}'
    order = {no: i for i, no in enumerate(seal_nos)}
    seals.sort(key=lambda s: order[s.seal_no])
    return seals, None


def _created_seal_payload(seal, operation=None):
    payload = {'seal': SealSerializer(seal).data}
    if operation is not None:
        payload['operation'] = lineage.operation_payload(operation)
    return payload


class SealListView(APIView):
    """封签列表与初始封装"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Seal.objects.select_related('goods').all()

        seal_no = request.query_params.get('seal_no')
        if seal_no:
            queryset = queryset.filter(seal_no__icontains=seal_no)
        batch_no = request.query_params.get('batch_no')
        if batch_no:
            queryset = queryset.filter(batch_no__icontains=batch_no)
        goods = request.query_params.get('goods')
        if goods:
            queryset = queryset.filter(goods_id=goods)
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        seals = queryset.order_by('-created_at', '-id')[start:start + page_size]
        return success_response(data={
            'list': SealSerializer(seals, many=True).data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """初始封装：建立谱系根节点"""
        data, err = _validate(SealCreateSerializer(data=request.data))
        if err:
            return err

        goods = Goods.objects.get(pk=data['goods'])
        try:
            seal, operation = lineage.seal(
                goods=goods,
                quantity=data['quantity'],
                seal_no=data.get('seal_no', ''),
                batch_no=data.get('batch_no', ''),
                operator=request.user,
                remark=data.get('remark', ''),
            )
        except LineageError as exc:
            return error_response(message=str(exc))

        logger.info(f"User {request.user.username} sealed {seal.seal_no} qty={seal.quantity}")
        return success_response(data=_created_seal_payload(seal, operation), message='封装成功')


class SealDetailView(APIView):
    """封签详情"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            seal = Seal.objects.select_related('goods').get(pk=pk)
        except Seal.DoesNotExist:
            return error_response(message='封签不存在', code=404)
        return success_response(data=SealSerializer(seal).data)


class SealSplitView(APIView):
    """拆分：一个大包装拆成多个小包装"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data, err = _validate(SplitSerializer(data=request.data))
        if err:
            return err

        try:
            parent = Seal.objects.get(seal_no=data['seal_no'])
        except Seal.DoesNotExist:
            return error_response(message=f'封签不存在：{data["seal_no"]}')

        try:
            operation, children = lineage.split(
                parent=parent,
                outputs=[
                    {'seal_no': o.get('seal_no', ''), 'batch_no': o.get('batch_no', ''),
                     'quantity': o['quantity']}
                    for o in data['outputs']
                ],
                operator=request.user,
                remark=data.get('remark', ''),
            )
        except LineageError as exc:
            return error_response(message=str(exc))

        logger.info(
            f"User {request.user.username} split {parent.seal_no} into "
            f"{[c.seal_no for c in children]}"
        )
        return success_response(data={
            'operation': lineage.operation_payload(operation),
            'seals': SealSerializer(children, many=True).data,
        }, message='拆分成功')


class SealMergeView(APIView):
    """合并：多个封签合并为一个新封签"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data, err = _validate(MergeSerializer(data=request.data))
        if err:
            return err

        seals, msg = _load_seals(data['seal_nos'])
        if msg:
            return error_response(message=msg)

        output = data['output']
        try:
            operation, child = lineage.merge(
                inputs=seals,
                output={
                    'seal_no': output.get('seal_no', ''),
                    'batch_no': output.get('batch_no', ''),
                    'quantity': output['quantity'],
                },
                operator=request.user,
                remark=data.get('remark', ''),
            )
        except LineageError as exc:
            return error_response(message=str(exc))

        logger.info(
            f"User {request.user.username} merged {[s.seal_no for s in seals]} into {child.seal_no}"
        )
        return success_response(
            data=_created_seal_payload(child, operation), message='合并成功'
        )


class SealRepackView(APIView):
    """重新封装：N 个投入重新封装为 M 个产出"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data, err = _validate(RepackSerializer(data=request.data))
        if err:
            return err

        seals, msg = _load_seals(data['seal_nos'])
        if msg:
            return error_response(message=msg)

        try:
            operation, children = lineage.repack(
                inputs=seals,
                outputs=[
                    {'seal_no': o.get('seal_no', ''), 'batch_no': o.get('batch_no', ''),
                     'quantity': o['quantity']}
                    for o in data['outputs']
                ],
                operator=request.user,
                remark=data.get('remark', ''),
            )
        except LineageError as exc:
            return error_response(message=str(exc))

        logger.info(
            f"User {request.user.username} repacked {[s.seal_no for s in seals]} into "
            f"{[c.seal_no for c in children]}"
        )
        return success_response(data={
            'operation': lineage.operation_payload(operation),
            'seals': SealSerializer(children, many=True).data,
        }, message='重新封装成功')


class SealIssueView(APIView):
    """领用：封签整体出库，离开可重组体系"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data, err = _validate(IssueSealsSerializer(data=request.data))
        if err:
            return err

        seals, msg = _load_seals(data['seal_nos'])
        if msg:
            return error_response(message=msg)

        try:
            operation = lineage.issue(
                seals=seals,
                receiver=data['receiver'],
                receiver_dept=data.get('receiver_dept', ''),
                operator=request.user,
                remark=data.get('remark', ''),
            )
        except LineageError as exc:
            return error_response(message=str(exc))

        logger.info(
            f"User {request.user.username} issued {[s.seal_no for s in seals]} "
            f"to {data['receiver']}"
        )
        return success_response(
            data={'operation': lineage.operation_payload(operation)}, message='领用成功'
        )


class SealFreezeView(APIView):
    """冻结：暂停封签的一切重组与领用"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data, err = _validate(FreezeSealsSerializer(data=request.data))
        if err:
            return err
        seals, msg = _load_seals(data['seal_nos'])
        if msg:
            return error_response(message=msg)
        try:
            operation = lineage.freeze(
                seals=seals, operator=request.user, remark=data.get('remark', '')
            )
        except LineageError as exc:
            return error_response(message=str(exc))
        logger.info(f"User {request.user.username} froze {data['seal_nos']}")
        return success_response(
            data={'operation': lineage.operation_payload(operation)}, message='冻结成功'
        )


class SealUnfreezeView(APIView):
    """解冻"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data, err = _validate(FreezeSealsSerializer(data=request.data))
        if err:
            return err
        seals, msg = _load_seals(data['seal_nos'])
        if msg:
            return error_response(message=msg)
        try:
            operation = lineage.unfreeze(
                seals=seals, operator=request.user, remark=data.get('remark', '')
            )
        except LineageError as exc:
            return error_response(message=str(exc))
        logger.info(f"User {request.user.username} unfroze {data['seal_nos']}")
        return success_response(
            data={'operation': lineage.operation_payload(operation)}, message='解冻成功'
        )


class SealLineageView(APIView):
    """谱系查询：任意封签的祖先、后代与全部操作"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            seal = Seal.objects.get(pk=pk)
        except Seal.DoesNotExist:
            return error_response(message='封签不存在', code=404)
        return success_response(data=lineage.get_lineage(seal))


class SealLineageByNoView(APIView):
    """按封签号查询谱系（审计现场通常只持有实物封签号）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, seal_no):
        try:
            seal = Seal.objects.get(seal_no=seal_no)
        except Seal.DoesNotExist:
            return error_response(message=f'封签不存在：{seal_no}', code=404)
        return success_response(data=lineage.get_lineage(seal))


def _parse_at_param(request):
    """解析 ?at= ISO 时间，返回 (datetime, None) 或 (None, 错误响应)。"""
    from django.utils.dateparse import parse_datetime
    from django.utils import timezone as dj_tz
    raw = request.query_params.get('at')
    at = parse_datetime(raw) if raw else None
    if raw and at is None:
        return None, error_response(
            message='at 参数须为 ISO 8601 时间，例如 2026-10-01T12:00:00+08:00'
        )
    if at is None:
        at = dj_tz.localtime()
    if dj_tz.is_naive(at):
        at = dj_tz.make_aware(at)
    return at, None


class SealTimelineView(APIView):
    """时点解释：历史任意时刻封签的状态与数量"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            seal = Seal.objects.get(pk=pk)
        except Seal.DoesNotExist:
            return error_response(message='封签不存在', code=404)
        at, err = _parse_at_param(request)
        if err:
            return err
        return success_response(data=lineage.explain_at(seal, at))


class SealTimelineByNoView(APIView):
    """按封签号解释历史时点状态"""
    def get(self, request, seal_no):
        try:
            seal = Seal.objects.get(seal_no=seal_no)
        except Seal.DoesNotExist:
            return error_response(message=f'封签不存在：{seal_no}', code=404)
        at, err = _parse_at_param(request)
        if err:
            return err
        return success_response(data=lineage.explain_at(seal, at))


class SealOperationListView(APIView):
    """操作台账：全部封签操作（只增不改的审计流水）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from .models import SealOperation
        queryset = SealOperation.objects.all()

        op_type = request.query_params.get('type')
        if op_type:
            queryset = queryset.filter(type=op_type)
        seal_no = request.query_params.get('seal_no')
        if seal_no:
            queryset = queryset.filter(models_in_seal_filter(seal_no)).distinct()

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        operations = queryset.order_by('-created_at', '-id')[start:start + page_size]
        return success_response(data={
            'list': [lineage.operation_payload(op) for op in operations],
            'total': total,
            'page': page,
            'page_size': page_size
        })


class SealConservationView(APIView):
    """守恒审计：全量复核每笔操作的守恒与状态约束"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        report = lineage.verify_conservation()
        return success_response(data=report,
                                message='谱系守恒校验通过' if report['healthy'] else '发现谱系异常')


def models_in_seal_filter(seal_no):
    """构造「操作涉及指定封签号」的 Q 条件。"""
    from django.db.models import Q
    return (
        Q(edges__source__seal_no=seal_no)
        | Q(edges__target__seal_no=seal_no)
        | Q(nodes__seal__seal_no=seal_no)
    )
